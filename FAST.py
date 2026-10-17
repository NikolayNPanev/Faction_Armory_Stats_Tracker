import json
import os
import re
from collections import Counter
from datetime import datetime
import openpyxl
from openpyxl.chart import LineChart, Reference

def get_faction_api_key(file_path):
    """Read the faction API key from a file."""
    try:
        with open(file_path, 'r') as f:
            api_key = f.read().strip()
            return api_key
    except FileNotFoundError:
        print(f"Error: The file {file_path} was not found.")
        return None
    except Exception as e:
        print(f"An error occurred while reading the API key: {e}")
        return None

def _format_cell_value(value):
    if value is None:
        return ""
    if isinstance(value, list):
        return ", ".join(str(v) for v in value)
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False)
    return value


def _loaned_to_rows(item_data, category=None):
    loaned_to_value = item_data.get("loaned_to")
    if loaned_to_value is None:
        return []
    if isinstance(loaned_to_value, str):
        values = [part.strip() for part in loaned_to_value.split(",") if part.strip()]
    elif isinstance(loaned_to_value, list):
        values = [str(v).strip() for v in loaned_to_value if str(v).strip()]
    else:
        values = [str(loaned_to_value)]

    counts = Counter(values)
    rows = []
    for loaned_to_id, count in counts.items():
        rows.append({
            "category": category,
            "ID": item_data.get("ID"),
            "name": item_data.get("name"),
            "loaned_to": loaned_to_id,
            "quantity": count,
        })
    return rows


def _is_item_list(value):
    return isinstance(value, list) and len(value) > 0 and isinstance(value[0], dict)


def _write_items_to_sheet(ws, items, title):
    ws.title = title
    headers = [h for h in list(items[0].keys()) if h != "loaned_to"]
    for col_num, header in enumerate(headers, 1):
        ws.cell(row=1, column=col_num, value=header)

    for row_num, row_data in enumerate(items, 2):
        for col_num, header in enumerate(headers, 1):
            ws.cell(row=row_num, column=col_num, value=_format_cell_value(row_data.get(header)))


def json_to_xlsx(json_file, xlsx_file):
    """Convert a JSON file to an XLSX file."""
    with open(json_file, 'r') as f:
        data = json.load(f)

    wb = openpyxl.Workbook()
    loaned_sheet = wb.active
    loaned_sheet.title = "Loaned To"
    loaned_headers = ["category", "ID", "name", "loaned_to", "quantity"]
    for col_num, header in enumerate(loaned_headers, 1):
        loaned_sheet.cell(row=1, column=col_num, value=header)

    loaned_row = 2
    category_sheets = {
        "temporary": "Temporary Items",
        "medical": "Medical Items",
        "drugs": "Drugs",
        "boosters": "Boosters",
        "armor": "Armor",
    }

    def process_items(items, title, category_name=None):
        nonlocal loaned_row
        if not _is_item_list(items):
            return
        if len(wb.sheetnames) == 1 and wb.sheetnames[0] == "Loaned To":
            ws = wb.active
            ws = wb.create_sheet(title=title)
        else:
            if title in wb.sheetnames:
                ws = wb[title]
            else:
                ws = wb.create_sheet(title=title)
        _write_items_to_sheet(ws, items, title)
        for row_data in items:
            for row in _loaned_to_rows(row_data, category_name):
                loaned_sheet.cell(row=loaned_row, column=1, value=_format_cell_value(row["category"]))
                loaned_sheet.cell(row=loaned_row, column=2, value=_format_cell_value(row["ID"]))
                loaned_sheet.cell(row=loaned_row, column=3, value=_format_cell_value(row["name"]))
                loaned_sheet.cell(row=loaned_row, column=4, value=_format_cell_value(row["loaned_to"]))
                loaned_sheet.cell(row=loaned_row, column=5, value=_format_cell_value(row["quantity"]))
                loaned_row += 1

    if isinstance(data, dict):
        if len(data) == 1 and _is_item_list(next(iter(data.values()))):
            first_key = next(iter(data.keys()))
            title = category_sheets.get(first_key, first_key.capitalize())
            process_items(next(iter(data.values())), title, first_key)
        else:
            for category_key, items in data.items():
                if category_key == "temporary":
                    category_title = "Temporary Items"
                elif category_key == "medical":
                    category_title = "Medical Items"
                elif category_key == "drugs":
                    category_title = "Drugs"
                elif category_key == "boosters":
                    category_title = "Boosters"
                elif category_key == "armor":
                    category_title = "Armor"
                else:
                    category_title = category_sheets.get(category_key, category_key.capitalize())
                process_items(items, category_title, category_key)
    elif _is_item_list(data):
        process_items(data, "Items")
    else:
        ws = wb.create_sheet(title="Data")
        headers = list(data.keys())
        for col_num, header in enumerate(headers, 1):
            ws.cell(row=1, column=col_num, value=header)
            ws.cell(row=2, column=col_num, value=_format_cell_value(data[header]))

    wb.save(xlsx_file)
    print(f"Successfully converted {json_file} to {xlsx_file}")

def torn_api_call_to_json(api_response, json_file):
    """Convert a Torn API response to a JSON file."""
    with open(json_file, 'w') as f:
        json.dump(api_response, f, indent=4)
    print(f"Successfully saved API response to {json_file}")


def analyze_item_usage_over_time(output_dir, report_file=None):
    """Compare timestamped JSON exports and write an XLSX usage report with charts."""
    pattern = re.compile(r"^(?P<category>.+)_(?P<timestamp>\d{8}_\d{6})\.json$")
    snapshots = []

    for filename in sorted(os.listdir(output_dir)):
        match = pattern.match(filename)
        if not match:
            continue

        category = match.group("category")
        timestamp_text = match.group("timestamp")
        timestamp = datetime.strptime(timestamp_text, "%Y%m%d_%H%M%S")
        json_path = os.path.join(output_dir, filename)

        with open(json_path, 'r') as handle:
            data = json.load(handle)

        item_quantities = {}
        if isinstance(data, dict):
            for _, value in data.items():
                if isinstance(value, list):
                    for item in value:
                        name = item.get("name") or item.get("ID") or "unknown"
                        item_quantities[name] = item.get("quantity", 0)
        elif isinstance(data, list):
            for item in data:
                name = item.get("name") or item.get("ID") or "unknown"
                item_quantities[name] = item.get("quantity", 0)

        snapshots.append({
            "category": category,
            "timestamp": timestamp_text,
            "datetime": timestamp,
            "item_quantities": item_quantities,
        })

    if not snapshots:
        return {"message": "No timestamped JSON exports found."}

    snapshots.sort(key=lambda item: item["datetime"])

    report_rows = [["category", "timestamp", "item", "quantity"]]
    item_keys = set()
    for snapshot in snapshots:
        for item_name in sorted(snapshot["item_quantities"].keys()):
            item_keys.add((snapshot["category"], item_name))
            report_rows.append([snapshot["category"], snapshot["timestamp"], item_name, snapshot["item_quantities"][item_name]])

    if report_file:
        wb = openpyxl.Workbook()
        summary_ws = wb.active
        summary_ws.title = "Usage Summary"
        for row in report_rows:
            summary_ws.append(row)
        summary_ws.freeze_panes = "A2"

        for column in summary_ws.columns:
            max_length = max(len(str(cell.value)) for cell in column if cell.value is not None)
            summary_ws.column_dimensions[column[0].column_letter].width = min(max_length + 2, 40)

        for category, item_name in sorted(item_keys):
            title = f"{category}_{item_name}".replace("/", "_")
            title = re.sub(r"[^A-Za-z0-9]+", "_", title).strip("_")
            title = title[:31]
            if title in wb.sheetnames:
                title = f"{title}_chart"
            chart_ws = wb.create_sheet(title=title)
            chart_ws.append(["timestamp", "quantity"])
            for snapshot in snapshots:
                if snapshot["category"] == category:
                    quantity = snapshot["item_quantities"].get(item_name, 0)
                    chart_ws.append([snapshot["timestamp"], quantity])
                else:
                    chart_ws.append([snapshot["timestamp"], 0])

            chart = LineChart()
            chart.type = "line"
            chart.style = 10
            chart.title = f"{category} - {item_name}"
            chart.y_axis.title = "Quantity"
            chart.x_axis.title = "Timestamp"

            data = Reference(chart_ws, min_row=2, max_row=chart_ws.max_row, min_col=2, max_col=2)
            categories = Reference(chart_ws, min_row=2, max_row=chart_ws.max_row, min_col=1, max_col=1)
            chart.add_data(data, titles_from_data=False)
            chart.set_categories(categories)
            chart.height = 7
            chart.width = 16
            chart_ws.add_chart(chart, "D2")
            chart_ws.cell(row=1, column=4, value="Chart")

        wb.save(report_file)
        print(f"Saved usage report to {report_file}")

    return report_rows


def torn_faction_api_call(api_key,attribute="basic"):
    """Make a Torn API call to get faction information."""
    import requests
    
    url = f"https://api.torn.com/faction/?selections={attribute}&key={api_key}"
    response = requests.get(url)
    
    if response.status_code == 200:
        return response.json()
    else:
        print(f"Error: Received status code {response.status_code} from the API.")
        return None

def torn_organized_crimes_api_call(api_key):
    """Retrieve the organized-crime catalogue and required items."""
    import requests

    url = f"https://api.torn.com/v2/torn/organizedcrimes?key={api_key}"
    try:
        response = requests.get(url, timeout=30)
        response.raise_for_status()
        data = response.json()
    except (requests.RequestException, json.JSONDecodeError) as error:
        print(f"Error retrieving organized-crime definitions: {error}")
        return None

    if not isinstance(data, dict) or not isinstance(data.get("organizedcrimes"), list):
        print(f"Error retrieving organized-crime definitions: {data}")
        return None
    return data

def torn_faction_inventory_api_call(api_key):
    """Retrieve every faction armory category, including loan status."""
    import requests

    categories = (
        "weapons", "armor", "temporary", "medical", "consumables",
        "drugs", "boosters", "utilities", "loot",
    )
    inventory = []
    inventory_timestamp = None
    page_size = 100
    for category in categories:
        offset = 0
        while True:
            try:
                response = requests.get(
                    "https://api.torn.com/v2/faction/inventory",
                    params={
                        "key": api_key,
                        "cat": category,
                        "limit": page_size,
                        "offset": offset,
                    },
                    timeout=30,
                )
                response.raise_for_status()
                data = response.json()
            except (requests.RequestException, json.JSONDecodeError) as error:
                print(f"Error retrieving faction {category} inventory: {error}")
                return None

            page = data.get("inventory") if isinstance(data, dict) else None
            if not isinstance(page, list):
                print(f"Error retrieving faction {category} inventory: {data}")
                return None

            inventory.extend(page)
            inventory_timestamp = data.get("inventory_timestamp", inventory_timestamp)
            if len(page) < page_size:
                break
            offset += len(page)

    return {"inventory": inventory, "inventory_timestamp": inventory_timestamp}
    
def main():
    api_key_file = 'faction_api_key.txt'
    api_key = get_faction_api_key(api_key_file)
    if api_key:
        output_dir = os.path.join(os.getcwd(), "OUTPUT")
        os.makedirs(output_dir, exist_ok=True)

        organized_crimes = torn_organized_crimes_api_call(api_key)
        if organized_crimes:
            definitions_file = os.path.join(output_dir, "organized_crime_definitions.json")
            with open(definitions_file, "w") as handle:
                json.dump(organized_crimes, handle, indent=4)
            print(f"Saved organized-crime definitions to: {definitions_file}")

        faction_inventory = torn_faction_inventory_api_call(api_key)
        if faction_inventory:
            inventory_file = os.path.join(output_dir, "faction_inventory.json")
            with open(inventory_file, "w") as handle:
                json.dump(faction_inventory, handle, indent=4)
            print(f"Saved faction inventory to: {inventory_file}")

        selections = {
            "temporary": "temporary",
            "medical": "medical",
            "drugs": "drugs",
            "boosters": "boosters",
            "armor": "armor",
        }

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

        for category, selection in selections.items():
            api_response = torn_faction_api_call(api_key, selection)
            if api_response:
                json_file = os.path.join(output_dir, f"{category}_{timestamp}.json")
                xlsx_file = os.path.join(output_dir, f"{category}_{timestamp}.xlsx")

                torn_api_call_to_json(api_response, json_file)
                json_to_xlsx(json_file, xlsx_file)
                print(f"Saved {category} files to: {output_dir}")
            else:
                print(f"Skipped {category} due to an API error.")

        report_file = os.path.join(output_dir, f"usage_report_{timestamp}.xlsx")
        usage_report = analyze_item_usage_over_time(output_dir, report_file)
        print(f"Usage comparison report generated with {len(usage_report) - 1} data rows.")
        print(f"All category files saved to: {output_dir}")

if __name__ == "__main__":
    main()